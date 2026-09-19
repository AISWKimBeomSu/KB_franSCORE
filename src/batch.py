"""일괄 조회 — 가맹점주 대출 신청 목록을 올리면 브랜드별 진단을 한 번에 붙인다.

왜 필요한가
    심사역은 브랜드를 하나씩 찾아보지 않는다. 하루에 들어온 가맹점주 대출 신청 수십 건을
    목록으로 받는다. 화면에서 브랜드를 한 건씩 검색해 결과를 옮겨 적게 하면 도구가 일을
    늘린다. 여기서는 목록(엑셀·CSV·붙여넣기)을 받아 브랜드를 찾아 붙이고, 등급·상태·
    중대 신호·확인 사항을 **같은 순서의 행**으로 돌려준다 — 받은 파일에 열이 붙어 나가는 형태.

매칭 원칙
    · 공시 등록명이 통칭과 다르므로(메가커피 → 메가엠지씨커피) 화면 검색과 **같은 규칙**
      (src/brand_search.py)을 쓴다. 정확 일치가 아니면 자동 확정하지 않고 '확인 필요'로 남긴다.
    · 같은 이름으로 등록된 브랜드가 둘 이상이면(국수나무 등) 가맹점이 많은 쪽을 붙이되
      '동명 브랜드'로 표시한다 — 잘못 붙은 진단이 조용히 결재에 올라가면 안 된다.
    · 평가 대상이 아닌 브랜드는 '공시에는 있으나 평가 대상 아님'과 '공시에 없음'을 구분한다.

차주 휴·폐업 (선택)
    목록에 사업자번호 열이 있으면 국세청 상태조회(src/nts.py)로 계속·휴업·폐업을 붙인다.
    공시는 1~2년 늦지만 국세청은 30분 주기로 갱신된다 — 신청인이 이미 폐업했다면 브랜드
    등급보다 먼저 봐야 할 사실이다. 키(DATA_GO_KR_KEY)가 없으면 '확인불가'와 설정 방법을 남긴다.

개인정보
    신청 목록에는 차주 정보가 섞여 올 수 있다. 이 모듈은 파일을 디스크에 쓰지 않고
    메모리에서만 처리한다. 브랜드·사업자번호 열 외의 열은 해석하지 않고 그대로 되돌려 준다.
"""
from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path

import pandas as pd

from src import grading, guidance
from src.brand_search import ALIASES, normalize, search

# 브랜드명이 들어 있을 법한 열 이름 (앞에 있을수록 우선)
BRAND_COLUMNS = ("브랜드명", "브랜드", "가맹브랜드", "영업표지", "가맹점 브랜드", "프랜차이즈",
                 "brand_name", "brand")
AMOUNT_HINTS = ("신청금액", "대출금액", "금액", "여신", "한도", "amount")
BNO_HINTS = ("사업자등록번호", "사업자번호", "b_no", "brno", "bizno")
BIZ_COLUMNS = ("사업자 상태", "폐업일", "사업자 확인")

MATCH_EXACT = "정확 일치"
MATCH_ALIAS = "통칭 일치"
MATCH_SAME_NAME = "동명 브랜드 — 확인 필요"
MATCH_FUZZY = "유사 일치 — 확인 필요"
MATCH_NOT_SCORED = "평가 대상 아님"
MATCH_NOT_FOUND = "공시에 없음"


@dataclass
class Context:
    """조회에 필요한 산출물 묶음 — 화면은 캐시된 것을 넘기고, 테스트는 파일에서 읽는다."""
    scores: pd.DataFrame
    findings: pd.DataFrame | None
    summary: pd.DataFrame | None
    all_brands: pd.DataFrame | None       # 공시 전체 브랜드명 (평가 대상 아님 판별용)
    outputs: Path

    @classmethod
    def from_files(cls, outputs: Path, processed: Path) -> Context:
        scores = pd.read_csv(outputs / "scores_latest.csv", encoding="utf-8-sig")
        fp, sp = outputs / "brand_diagnosis.parquet", outputs / "brand_diagnosis_summary.csv"
        pf = processed / "panel_full.parquet"
        return cls(
            scores=scores,
            findings=pd.read_parquet(fp) if fp.exists() else None,
            summary=pd.read_csv(sp, encoding="utf-8-sig") if sp.exists() else None,
            all_brands=(pd.read_parquet(pf, columns=["brand_name"]).drop_duplicates()
                        if pf.exists() else None),
            outputs=outputs)


def detect_brand_column(df: pd.DataFrame) -> str | None:
    cols = [str(c) for c in df.columns]
    for want in BRAND_COLUMNS:
        for c in cols:
            if c.strip().lower() == want.lower():
                return c
    for c in cols:                                   # '가맹 브랜드명(필수)' 같은 변형
        if "브랜드" in c or "영업표지" in c:
            return c
    return cols[0] if cols else None


def detect_amount_column(df: pd.DataFrame) -> str | None:
    for c in df.columns:
        if any(h in str(c) for h in AMOUNT_HINTS) and pd.to_numeric(df[c], errors="coerce").notna().any():
            return str(c)
    return None


def detect_bno_column(df: pd.DataFrame) -> str | None:
    """사업자번호 열 — 이름에 단서가 있고, 값의 절반 이상이 10자리 번호로 읽히는 열."""
    from src.nts import normalize_bno
    for c in df.columns:
        key = str(c).replace(" ", "").lower()
        if not any(h in key for h in BNO_HINTS):
            continue
        vals = df[c].dropna().astype(str).str.strip()
        vals = vals[vals != ""].head(50)
        if len(vals) and vals.map(lambda v: normalize_bno(v, check_digit=False) is not None).mean() >= 0.5:
            return str(c)
    return None


def attach_business_status(res: pd.DataFrame, bno_col: str, lookup=None) -> tuple[pd.DataFrame, dict]:
    """행마다 국세청 사업자 상태를 붙인다 — '사업자 상태'·'폐업일'·'사업자 확인' 세 열.

    번호를 적지 않은 행은 빈칸으로 둔다(확인불가와 구분). lookup 은 테스트 주입용이고,
    기본은 src.nts.lookup_status (키가 없으면 네트워크 없이 전 행 '확인불가' + 설정 방법).
    """
    from src import nts
    fn = lookup or nts.lookup_status
    raw = res[bno_col].astype(object).where(res[bno_col].astype(str).str.strip() != "", None)
    got = fn(list(raw)).reset_index(drop=True)
    blank = raw.isna().to_numpy()
    status = got["status"].astype(str).where(~blank, "")
    closed = pd.to_datetime(got["closed_on"], errors="coerce").dt.strftime("%Y-%m-%d").fillna("")
    note = got["error"].fillna("").astype(str)
    note = note.where(note != "", got["checksum_ok"].map(
        lambda ok: "" if ok else "검증번호 불일치 — 번호 오기 의심"))
    out = res.copy()
    pos = list(out.columns).index("매칭 상태") if "매칭 상태" in out.columns else len(out.columns)
    for i, (name, col) in enumerate(zip(BIZ_COLUMNS, (status, closed.where(~blank, ""),
                                                        note.where(~blank, "")), strict=True)):
        out.insert(pos + i, name, col.to_numpy())
    asked = ~blank
    summary = {"asked": int(asked.sum()),
               "closed": int((status == nts.STATUS_BY_CODE["03"]).sum()),
               "suspended": int((status == nts.STATUS_BY_CODE["02"]).sum()),
               "unknown": int(((status == nts.UNKNOWN) & asked).sum()),
               "no_key": bool(asked.any() and note[asked].str.contains(nts.KEY_ENV).all())}
    return out, summary


def _is_exact(query: str, name: str) -> bool:
    q = normalize(query)
    return normalize(name) in {q, normalize(ALIASES.get(q, "")) or q}


def match_one(query: str, ctx: Context) -> dict:
    """입력 브랜드명 하나 → {status, brand_id, brand_name, candidates}."""
    q = str(query or "").strip()
    if not q or q.lower() == "nan":
        return {"status": MATCH_NOT_FOUND, "brand_id": None, "brand_name": None, "candidates": ""}
    hit, near = search(ctx.scores, q, limit=5)
    if not hit.empty:
        top = hit.iloc[0]
        names = hit["brand_name"].astype(str)
        exact = hit[names.map(lambda n: _is_exact(q, n))]
        # 통칭 사전에 있는 이름("메가커피")이 등록명 하나에만 걸리면 확정해도 된다
        token = normalize(ALIASES.get(normalize(q), ""))
        alias = hit[names.map(lambda n: bool(token) and token in normalize(n))] if token else hit.head(0)
        if len(exact) == 1:
            status = MATCH_EXACT
        elif len(exact) > 1:
            status, top = MATCH_SAME_NAME, exact.iloc[0]
        elif len(alias) == 1:
            status, top = MATCH_ALIAS, alias.iloc[0]
        else:
            status = MATCH_FUZZY
        others = [str(n) for n in hit["brand_name"].astype(str) if n != str(top["brand_name"])][:3]
        return {"status": status, "brand_id": str(top["brand_id"]), "brand_name": str(top["brand_name"]),
                "candidates": " · ".join(others)}
    if ctx.all_brands is not None:
        h2, _ = search(ctx.all_brands, q, limit=3)
        if not h2.empty:
            return {"status": MATCH_NOT_SCORED, "brand_id": None, "brand_name": None,
                    "candidates": " · ".join(h2["brand_name"].astype(str).head(3))}
    return {"status": MATCH_NOT_FOUND, "brand_id": None, "brand_name": None,
            "candidates": " · ".join(near[:3])}


def screen(df: pd.DataFrame, brand_col: str, ctx: Context,
           amount_col: str | None = None) -> pd.DataFrame:
    """입력 표 → 입력 열 + 진단 열. 행 순서는 입력 그대로 둔다."""
    rates = grading.watch_rates(ctx.outputs)
    cuts = grading.cuts(ctx.outputs)
    scores = ctx.scores.assign(brand_id=ctx.scores["brand_id"].astype(str)).set_index("brand_id")
    fin = ctx.findings
    by_brand = ({str(k): g for k, g in fin.groupby(fin["brand_id"].astype(str))}
                if fin is not None and not fin.empty else {})
    summ = (ctx.summary.assign(brand_id=ctx.summary["brand_id"].astype(str)).set_index("brand_id")
            if ctx.summary is not None and not ctx.summary.empty else None)

    rows = []
    cache: dict[str, dict] = {}
    for q in df[brand_col].astype(str):
        m = cache.get(q) or match_one(q, ctx)
        cache[q] = m
        out = {"매칭 상태": m["status"], "매칭 브랜드": m["brand_name"] or "",
               "다른 후보": m["candidates"]}
        bid = m["brand_id"]
        if bid and bid in scores.index:
            r = scores.loc[bid]
            if isinstance(r, pd.DataFrame):
                r = r.iloc[0]
            g = guidance.grade_code(str(r.get("risk_grade")))
            f = by_brand.get(bid)
            risk = f[f["direction"] == "risk"] if f is not None else None
            crit = guidance.critical_findings(f)
            checks = guidance.checklist(risk if risk is not None and not risk.empty else None)
            docs = list(dict.fromkeys(d for c in checks for d in c["docs"]))
            head = (summ.loc[bid, "headline_detail"] if summ is not None and bid in summ.index else "")
            out.update({
                "업종": f"{r.get('industry_major', '')} · {r.get('industry_mid', '')}",
                "가맹점 수": pd.to_numeric(r.get("n_stores"), errors="coerce"),
                "등급": f"{g} {guidance.GRADE_LABEL.get(g, '')}" if g else "",
                "브랜드 상태": _state(r),
                "평가 경로": _basis(r),
                "브랜드 리스크(%)": grading.display_pct(r.get("deterioration_1y"), cuts),
                "1년 내 악화 위험(%)": round(grading.priority_risk(r, rates) * 100, 1),
                "중대 신호": " · ".join(str(x.get("title") or x.get("code")) for x in crit),
                "대표 소견": head if isinstance(head, str) else "",
                "확인 사항": "\n".join(f"☐ {c['check']}" for c in checks[:3]),
                "권고 확인 서류": " · ".join(docs[:3]),
                "권고 조치": guidance.grade_action(g),
            })
        rows.append(out)
    # 입력 열의 빈칸은 빈칸으로 — 엑셀의 빈 셀이 화면·반출에 'None' 글자로 찍히지 않게
    base = df.reset_index(drop=True)
    base = base.astype(object).where(base.notna(), "")
    res = pd.concat([base, pd.DataFrame(rows)], axis=1)
    if amount_col and amount_col in res.columns:
        res[amount_col] = pd.to_numeric(res[amount_col], errors="coerce")
    return res


def _basis(r) -> str:
    """정규 평가인지, 공시 공백 보정(src/coverage.py)으로 평가됐는지."""
    b = str(r.get("eligibility_basis") or "")
    return "정규" if not b or b in ("정규", "nan") else f"공시 공백 보정 — {b}"


def _state(r) -> str:
    s = str(r.get("brand_state") or "")
    if s == "요주의":
        k = pd.to_numeric(pd.Series([r.get("n_events_at_t")]), errors="coerce").iloc[0]
        return f"악화 발생 {int(k)}건" if pd.notna(k) else "악화 발생"
    return s or "-"


def summarize(res: pd.DataFrame, amount_col: str | None = None) -> dict:
    """결과 요약 — 화면 KPI 와 엑셀 '요약' 시트가 같은 값을 쓴다."""
    matched = res["매칭 브랜드"].astype(str).str.len() > 0
    need = res["매칭 상태"].isin([MATCH_SAME_NAME, MATCH_FUZZY])
    grade = res.get("등급", pd.Series("", index=res.index)).fillna("").astype(str)
    crit = res.get("중대 신호", pd.Series("", index=res.index)).fillna("").astype(str).str.len() > 0
    out = {"n": len(res), "matched": int(matched.sum()), "need_check": int(need.sum()),
           "unmatched": int((~matched).sum()), "fs3": int(grade.str.startswith("FS3").sum()),
           "critical": int(crit.sum()), "by_grade": []}
    amt = pd.to_numeric(res[amount_col], errors="coerce") if amount_col and amount_col in res else None
    for g in ("FS3 주의", "FS2 관찰", "FS1 안정", ""):
        m = grade == g
        if not m.any():
            continue
        out["by_grade"].append({"등급": g or "미평가", "건수": int(m.sum()),
                                "신청금액 합계": float(amt[m].sum()) if amt is not None else None})
    if amt is not None:
        out["amount_total"] = float(amt.sum())
        out["amount_fs3"] = float(amt[grade.str.startswith("FS3")].sum())
    return out


def to_excel(res: pd.DataFrame, summary: dict, meta: dict) -> bytes:
    """결과 · 요약 · 안내 세 시트. 결재에 첨부해도 설명 없이 읽히게 머리글은 한국어로."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        res.to_excel(w, index=False, sheet_name="조회결과")
        ws = w.sheets["조회결과"]
        ws.freeze_panes = "B2"
        wide = {"대표 소견": 60, "확인 사항": 70, "권고 확인 서류": 40, "권고 조치": 50, "업종": 18,
                "매칭 브랜드": 24, "다른 후보": 30, "중대 신호": 30}
        for i, col in enumerate(res.columns, start=1):
            letter = ws.cell(row=1, column=i).column_letter
            ws.column_dimensions[letter].width = wide.get(str(col), max(10, min(20, len(str(col)) * 2 + 2)))
        from openpyxl.styles import Alignment
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        pd.DataFrame(summary["by_grade"]).to_excel(w, index=False, sheet_name="요약")
        ws2 = w.sheets["요약"]
        base = len(summary["by_grade"]) + 3
        lines = [("전체 건수", summary["n"]), ("매칭 성공", summary["matched"]),
                 ("확인 필요(동명·유사)", summary["need_check"]),
                 ("미평가·미발견", summary["unmatched"]), ("중대 신호 보유", summary["critical"])]
        biz = summary.get("biz")
        if biz:
            lines += [("사업자 폐업(국세청)", biz["closed"]), ("사업자 휴업(국세청)", biz["suspended"]),
                      ("사업자 상태 확인불가", biz["unknown"])]
        for j, (k, v) in enumerate(lines):
            ws2.cell(row=base + j, column=1, value=k)
            ws2.cell(row=base + j, column=2, value=v)
        ws2.column_dimensions["A"].width = 22
        ws2.column_dimensions["B"].width = 16
        ws2.column_dimensions["C"].width = 16

        notes = [
            f"기준: {grading.year_label(meta.get('scored_year'))} (공정거래위원회 가맹사업 공시) · "
            f"산출 {meta.get('generated', '-')}",
            "등급은 브랜드의 구조악화 확률로 매긴 것이며 차주의 부도확률(PD)이 아닙니다.",
            "'악화 발생'은 올해 공시에 이미 악화 사건이 나타난 브랜드입니다. 이 구간의 위험은 모형이 아니라 "
            "같은 사건수 과거 브랜드의 다음 해 재발동 실현율입니다.",
            "'확인 필요' 행은 이름이 정확히 일치하지 않아 가장 가까운 브랜드를 붙인 것입니다 — 반드시 확인하십시오.",
            "이 자료는 2선 리스크 관리 참고용입니다. 여신 승인·거절, 한도·금리 결정에 사용하지 않습니다.",
            *(["'사업자 상태'는 국세청 사업자등록 상태조회(30분 주기 갱신) 결과입니다. 공시(연 1회)보다 "
               "최신이므로, 폐업·휴업 차주는 브랜드 등급과 무관하게 먼저 확인하십시오."]
              if summary.get("biz") else []),
        ]
        pd.DataFrame({"안내": notes}).to_excel(w, index=False, sheet_name="안내")
        w.sheets["안내"].column_dimensions["A"].width = 120
    return buf.getvalue()


def template_bytes() -> bytes:
    """업로드 양식 — 실제로 결과가 갈리는 예시를 담는다(정확·통칭·동명·평가 대상 아님)."""
    sample = pd.DataFrame({
        "신청번호": ["2026-0001", "2026-0002", "2026-0003", "2026-0004", "2026-0005", "2026-0006"],
        "브랜드명": ["메가커피", "인생냉면", "달리는커피", "빽다방", "국수나무", "크린토피아"],
        "신청금액(백만원)": [150, 80, 120, 200, 60, 180],
        "비고": ["통칭으로 입력해도 찾습니다", "", "", "", "같은 이름의 브랜드가 둘입니다",
                 "공시에는 있으나 외식 업종이 아니어서 평가 대상이 아닌 예"],
    })
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        sample.to_excel(w, index=False, sheet_name="신청목록")
        ws = w.sheets["신청목록"]
        for letter, width in zip("ABCD", (14, 20, 18, 34), strict=True):
            ws.column_dimensions[letter].width = width
    return buf.getvalue()


def read_upload(name: str, data: bytes) -> pd.DataFrame:
    """업로드 파일 → DataFrame. 엑셀은 첫 시트, CSV 는 UTF-8(BOM)·CP949 순서로 시도."""
    if name.lower().endswith((".xlsx", ".xlsm", ".xls")):
        return pd.read_excel(io.BytesIO(data), dtype=str)
    for enc in ("utf-8-sig", "cp949"):
        try:
            return pd.read_csv(io.BytesIO(data), dtype=str, encoding=enc)
        except UnicodeDecodeError:
            continue
    raise ValueError("CSV 인코딩을 읽지 못했습니다 (UTF-8 또는 CP949로 저장해 주십시오).")


def from_text(text: str) -> pd.DataFrame:
    """붙여넣기 — 한 줄에 브랜드 하나. 쉼표·탭으로 금액이 붙어 있으면 둘째 칸을 금액으로 본다."""
    rows = []
    for line in str(text or "").splitlines():
        parts = [p.strip() for p in line.replace("\t", ",").split(",")]
        if parts and parts[0]:
            rows.append({"브랜드명": parts[0], "신청금액": parts[1] if len(parts) > 1 else None})
    return pd.DataFrame(rows, columns=["브랜드명", "신청금액"])
