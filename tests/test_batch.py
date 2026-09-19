"""일괄 조회 — 예시 양식이 실제로 약속한 대로 갈리는지.

양식의 여섯 줄은 매칭 규칙을 하나씩 보여 주려고 고른 것이다: 통칭(메가커피), 정확 일치,
동명 브랜드(국수나무), 공시에는 있으나 평가 대상이 아닌 브랜드(크린토피아).
"""
from __future__ import annotations

import io
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from src import batch

ROOT = Path(__file__).resolve().parent.parent


@pytest.fixture(scope="module")
def ctx() -> batch.Context:
    return batch.Context.from_files(ROOT / "outputs", ROOT / "data" / "processed")


@pytest.fixture(scope="module")
def result(ctx) -> pd.DataFrame:
    df = pd.read_excel(io.BytesIO(batch.template_bytes()), dtype=str)
    return batch.screen(df, batch.detect_brand_column(df), ctx,
                        amount_col=batch.detect_amount_column(df))


def test_template_columns_are_detected():
    df = pd.read_excel(io.BytesIO(batch.template_bytes()), dtype=str)
    assert batch.detect_brand_column(df) == "브랜드명"
    assert batch.detect_amount_column(df) == "신청금액(백만원)"


def test_each_template_row_lands_in_the_promised_status(result):
    status = dict(zip(result["브랜드명"], result["매칭 상태"], strict=True))
    assert status["메가커피"] == batch.MATCH_ALIAS
    assert status["인생냉면"] == batch.MATCH_EXACT
    assert status["빽다방"] == batch.MATCH_EXACT
    assert status["국수나무"] == batch.MATCH_SAME_NAME
    assert status["크린토피아"] == batch.MATCH_NOT_SCORED


def test_input_order_and_columns_are_preserved(result):
    assert result["신청번호"].tolist() == [f"2026-000{i}" for i in range(1, 7)]
    for col in ("등급", "브랜드 상태", "중대 신호", "확인 사항", "권고 조치"):
        assert col in result.columns


def test_unscored_rows_carry_no_grade(result):
    row = result[result["브랜드명"] == "크린토피아"].iloc[0]
    assert not str(row.get("등급") or "").strip() or str(row.get("등급")) == "nan"


def test_summary_and_excel(result):
    s = batch.summarize(result, "신청금액(백만원)")
    assert s["n"] == 6 and s["unmatched"] == 1 and s["need_check"] == 1
    assert s["amount_total"] == pytest.approx(790.0)
    wb = load_workbook(io.BytesIO(batch.to_excel(result, s, {"scored_year": 2024, "generated": "t"})))
    assert wb.sheetnames == ["조회결과", "요약", "안내"]
    notes = " ".join(str(c.value) for c in wb["안내"]["A"])
    assert "부도확률" in notes and "승인" in notes


def test_pasted_text_parses_optional_amount():
    df = batch.from_text("메가커피, 150\n\n인생냉면\t80\n달리는커피")
    assert df["브랜드명"].tolist() == ["메가커피", "인생냉면", "달리는커피"]
    assert df["신청금액"].tolist()[:2] == ["150", "80"]


def test_csv_cp949_upload_is_read():
    raw = "브랜드명,금액\n인생냉면,80\n".encode("cp949")
    df = batch.read_upload("list.csv", raw)
    assert df.iloc[0]["브랜드명"] == "인생냉면"


# ── 차주 휴·폐업 (국세청 상태조회) ──────────────────────────────────────────

def test_business_number_column_is_detected_only_when_values_look_like_numbers():
    df = pd.DataFrame({"브랜드명": ["메가커피", "빽다방"], "사업자등록번호": ["124-81-00998", "2208162517"],
                       "사업자 구분": ["개인", "법인"]})
    assert batch.detect_bno_column(df) == "사업자등록번호"
    assert batch.detect_bno_column(df.drop(columns="사업자등록번호")) is None


def test_business_status_columns_are_attached_before_match_columns(ctx):
    df = pd.DataFrame({"브랜드명": ["메가커피", "빽다방", "국수나무", "교촌치킨"],
                       "사업자번호": ["124-81-00998", "", "220-81-62517", "12345"]})
    res = batch.screen(df, "브랜드명", ctx)

    def fake_lookup(values):
        assert values[1] is None                     # 빈칸은 묻지 않는다
        return pd.DataFrame({
            "input": [str(v) for v in values], "b_no": ["1248100998", None, "2208162517", None],
            "status": ["폐업", "확인불가", "휴업", "확인불가"], "status_code": ["03", "", "02", ""],
            "closed_on": [pd.Timestamp("2026-03-31"), pd.NaT, pd.NaT, pd.NaT], "tax_type": ["", "", "", ""],
            "error": [None, "형식 오류 — 숫자 10자리가 아닙니다.", None, "형식 오류 — 숫자 10자리가 아닙니다."],
            "checksum_ok": [True, False, True, False]})

    out, s = batch.attach_business_status(res, "사업자번호", lookup=fake_lookup)
    cols = list(out.columns)
    assert cols.index("사업자 상태") < cols.index("매칭 상태")
    assert out["사업자 상태"].tolist() == ["폐업", "", "휴업", "확인불가"]
    assert out["폐업일"].tolist() == ["2026-03-31", "", "", ""]
    assert "형식 오류" in out.loc[3, "사업자 확인"] and out.loc[1, "사업자 확인"] == ""
    assert s == {"asked": 3, "closed": 1, "suspended": 1, "unknown": 1, "no_key": False}


def test_without_key_rows_say_how_to_enable_and_no_network(ctx, monkeypatch):
    import requests

    from src import nts
    monkeypatch.delenv(nts.KEY_ENV, raising=False)
    monkeypatch.setattr(nts, "load_secrets", lambda: [])
    monkeypatch.setattr(requests, "post", lambda *a, **k: (_ for _ in ()).throw(AssertionError("network")))
    res = batch.screen(pd.DataFrame({"브랜드명": ["메가커피", "빽다방"], "사업자번호": ["124-81-00998", "12345"]}),
                       "브랜드명", ctx)
    out, s = batch.attach_business_status(res, "사업자번호")
    assert out.loc[0, "사업자 상태"] == nts.UNKNOWN and nts.DATASET_ID in out.loc[0, "사업자 확인"]
    assert "형식 오류" in out.loc[1, "사업자 확인"]
    assert s["no_key"], "형식 오류 행이 섞여도 키 안내는 나와야 한다"
